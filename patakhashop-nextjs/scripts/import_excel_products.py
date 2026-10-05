import openpyxl
import os
import json
import urllib.request
import urllib.error

# Load Excel File
excel_path = r'c:\phatka shop\Product_Catalogue_With_Images_Completed (1).xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Catalogue']

# Check existing image files in public/products
images_dir = r'c:\phatka shop\patakhashop-nextjs\public\products'
os.makedirs(images_dir, exist_ok=True)
existing_files = os.listdir(images_dir)

row_image_filename = {}
for fname in existing_files:
    if fname.startswith('cat_img_row_'):
        parts = fname.split('_')
        row_num = int(parts[3].split('.')[0])
        row_image_filename[row_num] = f'/products/{fname}'

category_map = {
    '2&3 sound': 'crackers',
    'atom fountain': 'flowerpots',
    'bomb': 'bombs',
    'fancy': 'skyshots',
    'fancyfan': 'skyshots',
    'flower pot': 'flowerpots',
    'gc': 'chakkar',
    'novelties': 'kids special',
    'pencil': 'torches',
    'rocket': 'rockets',
    'sparkler': 'sparklers'
}

products_json = []
products_db = []
sno = 1

for r in range(2, ws.max_row + 1):
    brand = str(ws.cell(r, 1).value or '').strip()
    prod_name = str(ws.cell(r, 2).value or '').strip()
    cat_raw = str(ws.cell(r, 3).value or '').strip()
    sub_cat = str(ws.cell(r, 4).value or '').strip()
    prod_type = str(ws.cell(r, 5).value or '').strip()
    web_name = str(ws.cell(r, 6).value or '').strip()
    desc = str(ws.cell(r, 10).value or '').strip()
    
    name = web_name if (web_name and web_name != 'None') else prod_name
    if not name or name == 'None':
        continue
        
    cat_lower = cat_raw.lower()
    category = category_map.get(cat_lower, 'all')
    
    tags_list = [t for t in [cat_raw, sub_cat, prod_type] if t and t.lower() != 'none']
    tags = ', '.join(tags_list) if tags_list else category
    
    image_url = row_image_filename.get(r, '/logo-62.jpg')
    
    item_id = str(sno)
    
    pj = {
        'id': item_id,
        'sno': sno,
        'name': name,
        'brand': brand if (brand and brand != 'None') else 'Patakha Shop',
        'category': category,
        'tags': tags,
        'price': 0,
        'originalPrice': 0,
        'discount': 'Price Pending',
        'inStock': True,
        'image': image_url,
        'stockQuantity': 50
    }
    products_json.append(pj)
    
    pdb = {
        'id': item_id,
        'sno': sno,
        'name': name,
        'brand': brand if (brand and brand != 'None') else 'Patakha Shop',
        'category': category,
        'tags': tags,
        'price': 0,
        'original_price': 0,
        'discount': 'Price Pending',
        'in_stock': True,
        'stock_quantity': 50,
        'image': image_url
    }
    products_db.append(pdb)
    sno += 1

# Write to products.json
json_path = r'c:\phatka shop\patakhashop-nextjs\src\data\products.json'
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(products_json, f, indent=2, ensure_ascii=False)
print(f'Wrote {len(products_json)} products to {json_path}')

# Write seed_products.sql
sql_lines = [
    '-- Supabase Seed Data for Products (Generated from Excel Master Catalogue)',
    'DELETE FROM products;',
    'INSERT INTO products (id, sno, name, brand, category, tags, price, original_price, discount, in_stock, stock_quantity, image) VALUES'
]

vals = []
for p in products_db:
    esc_name = p['name'].replace("'", "''")
    esc_brand = p['brand'].replace("'", "''")
    esc_cat = p['category'].replace("'", "''")
    esc_tags = p['tags'].replace("'", "''")
    esc_disc = p['discount'].replace("'", "''")
    esc_img = p['image'].replace("'", "''")
    
    v = f"('{p['id']}', {p['sno']}, '{esc_name}', '{esc_brand}', '{esc_cat}', '{esc_tags}', 0, 0, '{esc_disc}', true, 50, '{esc_img}')"
    vals.append(v)

sql_lines.append(',\n'.join(vals) + ';')

sql_path = r'c:\phatka shop\patakhashop-nextjs\supabase\seed_products.sql'
with open(sql_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(sql_lines))
print(f'Wrote SQL seed file to {sql_path}')

# Write to memoryStore initial state or updated backup
backup_path = r'c:\phatka shop\patakhashop-nextjs\src\data\products_backup_full.json'
with open(backup_path, 'w', encoding='utf-8') as f:
    json.dump(products_json, f, indent=2, ensure_ascii=False)

print('Script completed successfully.')
